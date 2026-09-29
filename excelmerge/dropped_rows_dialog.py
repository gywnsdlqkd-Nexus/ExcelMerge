# -*- coding: utf-8 -*-
"""'비교에서 빠진 행' 진단 창 — 왜 빠졌는지 + 어느 열을 키로 쓰면 되는지.

키 열에 중복·빈 값이 있으면 그 행들은 비교 대상에서 빠진다(compute_diff 는 같은 키의
첫 행만 쓴다). 지금까지는 상태바 한 줄로만 알려서, 실제로 14,659행이 빠진 비교가
'차이 없음'처럼 보이는 일이 있었다.

실측해 보면 현장 데이터의 원인은 거의 전부 **키 열을 잘못 잡은 것**이다
(Data_SkillModule_CS.xlsx 를 B열 키로 비교 → 14,659행 전부 중복 키, 같은 시트의
A열(ModuleID)은 완벽히 유니크). 그래서 이 창의 중심은 목록이 아니라
**'이 열을 키로 바꾸기'** 버튼이다.
"""
from PyQt5.QtCore import Qt
from PyQt5.QtWidgets import (
    QDialog, QVBoxLayout, QHBoxLayout, QLabel, QPushButton, QTableWidget,
    QTableWidgetItem, QAbstractItemView, QHeaderView,
)
from openpyxl.utils import get_column_letter

from .theme import ui_font


class DroppedRowsDialog(QDialog):
    """info = diff_engine.dropped_key_rows(...), candidates = unique_key_candidates(...)."""

    def __init__(self, parent, key_col: int, info: dict, candidates: list,
                 total: int, on_pick_key=None):
        super().__init__(parent)
        self.setWindowTitle("비교에서 빠진 행")
        self.resize(720, 560)
        self._on_pick_key = on_pick_key
        lay = QVBoxLayout(self)

        key_name = (f"{get_column_letter(key_col + 1)}열"
                    if key_col is not None and key_col >= 0 else "키 없음")
        head = QLabel(
            f"<b>키 열: {key_name}</b><br>"
            f"A 파일 {info['per_side']['A']:,}행 · B 파일 {info['per_side']['B']:,}행이 "
            f"비교에서 빠졌습니다 (합계 {total:,}행)<br>"
            f"· 키 중복 {info['dup']:,}　· 빈 키 {info['blank']:,}")
        head.setFont(ui_font(10))
        head.setTextFormat(Qt.RichText)
        lay.addWidget(head)

        if info["dup_keys"]:
            lay.addWidget(self._section("가장 많이 겹친 키"))
            top = QTableWidget(len(info["dup_keys"]), 3, self)
            top.setHorizontalHeaderLabels(["키", "A 중복", "B 중복"])
            for r, (k, na, nb) in enumerate(info["dup_keys"]):
                top.setItem(r, 0, QTableWidgetItem(str(k)))
                top.setItem(r, 1, QTableWidgetItem(f"{na:,}행"))
                top.setItem(r, 2, QTableWidgetItem(f"{nb:,}행"))
            self._tidy(top, max_height=170)
            lay.addWidget(top)

        lay.addWidget(self._section("키로 쓸 수 있는 열 (양쪽 파일 모두 유니크·빈 값 없음)"))
        if candidates:
            for col, header in candidates[:5]:
                row = QHBoxLayout()
                label = QLabel(f"● {get_column_letter(col + 1)}열"
                               + (f" ({header})" if header else ""))
                label.setFont(ui_font(10))
                row.addWidget(label)
                row.addStretch()
                btn = QPushButton("이 열을 키로 바꾸기")
                btn.setFont(ui_font(9))
                btn.setFocusPolicy(Qt.NoFocus)
                btn.clicked.connect(lambda _=False, c=col: self._pick(c))
                row.addWidget(btn)
                lay.addLayout(row)
        else:
            none = QLabel("이 시트에는 유니크한 열이 없습니다 — 키 열 해제(ROW 순서 비교)를 "
                          "쓰거나, 중복이 정상인 데이터인지 확인하세요.")
            none.setWordWrap(True)
            none.setFont(ui_font(9))
            lay.addWidget(none)

        shown = len(info["rows"])
        lay.addWidget(self._section(
            f"빠진 행 목록 — 총 {total:,}행 중 {shown:,}행 표시"
            + ("  (Ctrl+C 로 복사)" if shown else "")))
        table = QTableWidget(shown, 4, self)
        table.setHorizontalHeaderLabels(["파일", "원본 행", "키", "사유"])
        for r, (side, row_no, key, reason) in enumerate(info["rows"]):
            table.setItem(r, 0, QTableWidgetItem(side))
            table.setItem(r, 1, QTableWidgetItem(f"{row_no:,}"))
            table.setItem(r, 2, QTableWidgetItem(str(key)))
            table.setItem(r, 3, QTableWidgetItem(reason))
        self._tidy(table)
        lay.addWidget(table, 1)

        close = QPushButton("닫기")
        close.setFocusPolicy(Qt.NoFocus)
        close.clicked.connect(self.reject)
        bottom = QHBoxLayout()
        bottom.addStretch()
        bottom.addWidget(close)
        lay.addLayout(bottom)

    @staticmethod
    def _section(text: str) -> QLabel:
        lbl = QLabel(f"▸ {text}")
        lbl.setFont(ui_font(9, bold=True))
        return lbl

    @staticmethod
    def _tidy(table: QTableWidget, max_height: int = 0):
        table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        table.setSelectionBehavior(QAbstractItemView.SelectRows)
        table.verticalHeader().setVisible(False)
        table.setFont(ui_font(9))
        table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents)
        table.horizontalHeader().setStretchLastSection(True)
        if max_height:
            table.setMaximumHeight(max_height)

    def _pick(self, col: int):
        """키 열 교체는 기존 경로(key_col_changed)를 그대로 탄다 — 새 경로를 만들지 않는다."""
        if self._on_pick_key is not None:
            self._on_pick_key(col)
        self.accept()

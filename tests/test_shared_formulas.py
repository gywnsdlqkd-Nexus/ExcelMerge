# -*- coding: utf-8 -*-
"""공유 수식이 저장 뒤에도 성립하는가 — Excel 이 파일을 거부하지 않도록.

엑셀은 같은 모양의 수식을 한 번만 적는다.

    <c r="B4"><f t="shared" ref="B4:B16" si="0">B3+1</f><v>103</v></c>   주인
    <c r="B5"><f t="shared" si="0"/><v>104</v></c>                       추종자

추종자는 수식 본문이 없다. 주인을 보고 자기 자리에 맞게 옮겨 쓴다. 그래서 주인이
망가지면 추종자는 **읽을 수 없는 셀**이 되고 Excel 은 "내용에 문제가 있습니다" 복구
창을 띄운다.

실기에서 두 파일이 이 상태로 나갔다(40_Build). 원인은 둘이었다:

  1. **주인이 사라졌다** — 병합이 주인 셀을 값으로 덮었다.
     (Data_MiniGameRoulette_CS.xlsx: si=2 추종자 3개에 주인이 없었다)
  2. **주인의 ref 가 자기를 안 담는다** — 행을 지우면 <c r> 은 당겨지는데 ref 는
     그대로 남았다. (Data_TextUITable_CS.xlsm: 주인 A5250 의 ref 가 A5251:A5314)

둘 다 **값 대조로는 안 잡힌다.** <v> 의 캐시 값이 그대로라 셀은 전부 일치한다. 그래서
저장 검증에 구조 검사를 넣었고, 아래 '가드' 절이 그게 실제로 막는지를 확인한다.
"""
import os
import zipfile

import openpyxl
import pytest
from lxml import etree

from excelmerge.loaders import clear_values_cache, load_values_any
from excelmerge.xlsx_writer import (_bounding_ref, _ref_in,
                                    _formula_mismatches,
                                    _write_patches_to_file)

NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"

# A        B(공유수식 그룹)   C
# 1 head   head              head
# 2 r2     100               x2
# 3 r3     =B2+1  주인       x3      ref B3:B6, si=0
# 4 r4     추종자            x4
# 5 r5     추종자            x5
# 6 r6     추종자            x6
SHEET = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
<dimension ref="A1:C6"/><sheetData>
<row r="1"><c r="A1" t="str"><v>ID</v></c><c r="B1" t="str"><v>N</v></c><c r="C1" t="str"><v>M</v></c></row>
<row r="2"><c r="A2" t="str"><v>r2</v></c><c r="B2"><v>100</v></c><c r="C2" t="str"><v>x2</v></c></row>
<row r="3"><c r="A3" t="str"><v>r3</v></c><c r="B3"><f t="shared" ref="B3:B6" si="0">B2+1</f><v>101</v></c><c r="C3" t="str"><v>x3</v></c></row>
<row r="4"><c r="A4" t="str"><v>r4</v></c><c r="B4"><f t="shared" si="0"/><v>102</v></c><c r="C4" t="str"><v>x4</v></c></row>
<row r="5"><c r="A5" t="str"><v>r5</v></c><c r="B5"><f t="shared" si="0"/><v>103</v></c><c r="C5" t="str"><v>x5</v></c></row>
<row r="6"><c r="A6" t="str"><v>r6</v></c><c r="B6"><f t="shared" si="0"/><v>104</v></c><c r="C6" t="str"><v>x6</v></c></row>
</sheetData></worksheet>"""


@pytest.fixture
def book(tmp_path):
    """공유 수식이 든 xlsx. openpyxl 은 공유 수식을 만들지 못해 XML 을 직접 넣는다."""
    p = tmp_path / "a.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(1, 7):
        ws.append(["", "", ""])
    wb.save(str(p))

    tmp = str(p) + ".build"
    with zipfile.ZipFile(str(p)) as zin, \
         zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in zin.namelist():
            data = SHEET.encode("utf-8") if n.endswith("worksheets/sheet1.xml") \
                else zin.read(n)
            zout.writestr(n, data)
    os.replace(tmp, str(p))
    assert _formula_mismatches(str(p)) == [], "전제: 저장 전은 성립한다"
    clear_values_cache()
    return p


def _formulas(path):
    """{셀: (t, ref, si, 수식본문)} — <f> 가 있는 셀만."""
    out = {}
    with zipfile.ZipFile(str(path)) as z:
        root = etree.fromstring(z.read("xl/worksheets/sheet1.xml"))
        for c in root.iter(NS + "c"):
            f = c.find(NS + "f")
            if f is not None:
                out[c.get("r")] = (f.get("t"), f.get("ref"), f.get("si"),
                                   (f.text or "").strip())
    return out


def _save(path, **kw):
    clear_values_cache()
    _write_patches_to_file(str(path), kw.pop("patches", {}), **kw)


def _values(path):
    clear_values_cache()
    return load_values_any(str(path))


# ── 1. 주인이 사라지는 네 가지 경로 ─────────────────────────────────────────
# 저장이 주인을 망가뜨릴 수 있는 길은 넷이다. 하나씩 땜질하지 않고 바꾼 **뒤에**
# 한 번에 되돌리므로, 넷 다 같은 규칙으로 지켜져야 한다.

@pytest.mark.parametrize("kw, why", [
    ({"patches": {"B3": "555"}},                 "주인 셀을 값으로 덮어쓴다"),
    ({"patches": {"B3": ""}},                    "주인 셀을 빈값으로 지운다"),
    ({"delete_row_nums": {3}},                   "주인이 든 행을 지운다"),
    ({"delete_row_nums": {2}},                   "주인 위의 행을 지운다"),
    ({"delete_col_letters": {"B"}},              "그룹이 든 열을 지운다"),
    ({"patches": {"B3": "555"}, "delete_row_nums": {2}}, "덮어쓰기와 행 삭제를 같이"),
])
def test_the_saved_file_keeps_its_shared_formulas_valid(book, kw, why):
    """이 테스트가 그 복구 창이다 — 실패하면 Excel 이 파일을 거부한다."""
    _save(book, **kw)
    assert _formula_mismatches(str(book)) == [], why


# ── 2. 주인이 사라지면 남은 셀을 주인으로 올린다 ────────────────────────────

def test_a_survivor_is_promoted_to_master(book):
    _save(book, patches={"B3": "555"})
    f = _formulas(book)
    assert "B3" not in f, "덮어쓴 셀에 수식이 남았다"
    assert f["B4"][1], f"새 주인에 ref 가 없다: {f}"
    assert f["B4"][0] == "shared" and f["B4"][2] == "0"


def test_the_promoted_formula_is_moved_to_its_new_place(book):
    """수식 본문은 주인 자리 기준이다 — 새 자리로 옮겨 쓰지 않으면 값이 틀린다.

    B3 의 'B2+1' 이 B4 로 올라가면 'B3+1' 이 돼야 한다.
    """
    _save(book, patches={"B3": "555"})
    assert _formulas(book)["B4"][3] == "B3+1", _formulas(book)


def test_the_promoted_master_covers_the_rest(book):
    _save(book, patches={"B3": "555"})
    _, ref, _, _ = _formulas(book)["B4"]
    for cell in ("B4", "B5", "B6"):
        assert _ref_in(cell, ref), f"{cell} 이 ref={ref} 밖이다"


def test_deleting_the_master_row_promotes_too(book):
    """주인 행 자체가 사라지는 경우 — 번호가 당겨진 뒤에도 성립해야 한다."""
    _save(book, delete_row_nums={3})
    assert _formula_mismatches(str(book)) == []
    f = _formulas(book)
    masters = [r for r, (t, ref, _, _) in f.items() if t == "shared" and ref]
    assert len(masters) == 1, f"주인이 {len(masters)}개: {f}"


# ── 3. 주인이 살아남으면 ref 만 바로잡는다 ──────────────────────────────────

def test_the_ref_follows_the_renumbered_rows(book):
    """행을 지우면 <c r> 은 당겨진다. ref 가 안 따라가면 주인이 자기 범위 밖이 된다 —
    실기에서 Data_TextUITable_CS.xlsm 이 바로 이 상태였다.

    1행(헤더)을 지운다. 주인 B3 은 B2 로 당겨지지만 그 참조(B2)는 지워지지 않아
    그룹이 살아남는다 — ref 가 따라오는지만 깨끗하게 볼 수 있다.
    """
    _save(book, delete_row_nums={1})
    f = _formulas(book)
    own = next(r for r, (t, ref, _, _) in f.items() if t == "shared" and ref)
    assert own == "B2", f"주인이 B2 로 당겨져야 한다: {f}"
    assert _ref_in(own, f[own][1]), f"주인 {own} 이 자기 ref={f[own][1]} 밖이다"


def test_the_master_body_follows_the_renumbered_rows_too(book):
    """ref 뿐 아니라 **수식 본문**도 따라가야 한다.

    주인 B3 의 'B2+1' 은 1행이 사라지면 B2 를 가리킬 수 없다 — 그 셀이 B1 이 됐다.
    본문을 그대로 두면 B2(=자기 자신)를 가리켜 순환 참조가 된다.
    """
    _save(book, delete_row_nums={1})
    f = _formulas(book)
    own = next(r for r, (t, ref, _, _) in f.items() if t == "shared" and ref)
    assert f[own][3] == "B1+1", f"본문이 안 따라왔다: {f}"


def test_a_master_pointing_at_a_deleted_row_becomes_ref_error(book):
    """지워진 행을 가리키면 #REF! — 엑셀이 하는 그대로다.

    주인 B3 의 'B2+1' 에서 2행을 지우면 가리킬 셀이 없어진다. 그룹은 더 못 쓰므로
    각 칸이 제 수식을 갖게 되고(엑셀도 깨진 칸을 떼어낸다), 값은 그대로 남는다.
    """
    _save(book, delete_row_nums={2})
    f = _formulas(book)
    assert any("#REF!" in body for *_x, body in f.values()), f
    assert _formula_mismatches(str(book)) == []


# ── 4. 혼자 남으면 공유할 이유가 없다 ───────────────────────────────────────

def test_a_lone_survivor_becomes_a_plain_formula(book):
    """추종자가 다 사라지면 평범한 수식으로 되돌린다 — 혼자인 공유 수식은 군더더기다."""
    _save(book, delete_row_nums={4, 5, 6})
    f = _formulas(book)
    assert len(f) == 1, f
    t, ref, si, body = next(iter(f.values()))
    assert t is None and ref is None and si is None, f"공유 표시가 남았다: {f}"
    assert body == "B2+1"


# ── 5. 값은 건드리지 않는다 ─────────────────────────────────────────────────

def test_the_values_are_untouched(book):
    before = _values(book)
    _save(book, patches={"B3": "555"})
    after = _values(book)
    assert after[0] == before[0]
    assert after[2][1] == "555", after[2]
    for r in (3, 4, 5):                    # 나머지 캐시 값은 그대로
        assert after[r][1] == before[r][1], f"{r+1}행 값이 바뀌었다"


def test_a_file_without_shared_formulas_is_unaffected(tmp_path):
    p = tmp_path / "plain.xlsx"
    wb = openpyxl.Workbook()
    for r in (["ID", "N"], ["k1", "1"], ["k2", "2"]):
        wb.active.append(r)
    wb.save(str(p))
    _save(p, patches={"B2": "9"})
    assert _formula_mismatches(str(p)) == []
    assert _values(p)[1][1] == "9"


# ── 6. 가드가 살아 있나 ─────────────────────────────────────────────────────
# 값만 보는 검증이 이 사고를 통과시켰다. 구조 검사를 **넣었다**고 끝이 아니라,
# 그게 실제로 막는지를 봐야 한다 — 복원을 꺼 놓고 저장을 시도한다.

def test_the_verifier_blocks_a_broken_shared_formula(book, monkeypatch):
    import excelmerge.xlsx_writer as W
    monkeypatch.setattr(W, "_repair_shared_formulas", lambda *a, **k: None)
    before = _values(book)
    with pytest.raises(ValueError) as e:
        _save(book, patches={"B3": "555"})
    assert "공유수식" in str(e.value), str(e.value)
    assert _values(book) == before, "원본이 덮였다"
    assert not os.path.exists(str(book) + ".tmp_merge"), "임시 파일이 남았다"


def test_the_verifier_catches_the_self_excluding_ref(book, monkeypatch):
    """실기 원인 2 — 수식 ref 가 자기 셀을 안 담는다.

    행 삭제 때 ref 를 당기는 일(_shift_ref_rows)과 공유 그룹 복원, 둘 다 꺼야 이
    상태가 재현된다. 둘 중 하나만 살아 있어도 막히므로 안전장치가 겹쳐 있는 셈이다.
    """
    import excelmerge.xlsx_writer as W
    monkeypatch.setattr(W, "_shift_ref_rows", lambda ref, deleted: ref)
    monkeypatch.setattr(W, "_repair_shared_formulas", lambda *a, **k: None)
    with pytest.raises(ValueError) as e:
        _save(book, delete_row_nums={2})
    assert "ref" in str(e.value) and "담지 못합니다" in str(e.value), str(e.value)


# ── 8. 배열 수식 — ref 를 갖는 건 공유 수식만이 아니다 ─────────────────────
# 동적 배열 수식도 <f t="array" ref="H117"> 처럼 자기 자리를 ref 에 적는다. 셀만
# 당기고 ref 를 두면 Excel 이 파일을 거부한다(실측: Data_Illustration_C.xlsx).

ARRAY_SHEET = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
<dimension ref="A1:B4"/><sheetData>
<row r="1"><c r="A1" t="str"><v>ID</v></c><c r="B1" t="str"><v>N</v></c></row>
<row r="2"><c r="A2" t="str"><v>r2</v></c><c r="B2"><v>7</v></c></row>
<row r="3"><c r="A3" t="str"><v>r3</v></c><c r="B3" cm="1"><f t="array" ref="B3">SUM(B2)</f><v>7</v></c></row>
<row r="4"><c r="A4" t="str"><v>r4</v></c><c r="B4"><v>9</v></c></row>
</sheetData></worksheet>"""


@pytest.fixture
def array_book(tmp_path):
    p = tmp_path / "arr.xlsx"
    wb = openpyxl.Workbook()
    for _ in range(4):
        wb.active.append(["", ""])
    wb.save(str(p))
    tmp = str(p) + ".build"
    with zipfile.ZipFile(str(p)) as zin, \
         zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in zin.namelist():
            zout.writestr(n, ARRAY_SHEET.encode("utf-8")
                          if n.endswith("worksheets/sheet1.xml") else zin.read(n))
    os.replace(tmp, str(p))
    assert _formula_mismatches(str(p)) == [], "전제: 저장 전은 성립한다"
    clear_values_cache()
    return p


def test_an_array_ref_follows_the_renumbered_rows(array_book):
    _save(array_book, delete_row_nums={2})
    f = _formulas(array_book)
    assert "B2" in f, f"배열 수식이 B2 로 당겨져야 한다: {f}"
    assert f["B2"][1] == "B2", f"ref 가 안 따라왔다: {f}"
    assert _formula_mismatches(str(array_book)) == []


def test_the_verifier_catches_a_stale_array_ref(array_book, monkeypatch):
    import excelmerge.xlsx_writer as W
    monkeypatch.setattr(W, "_shift_ref_rows", lambda ref, deleted: ref)
    with pytest.raises(ValueError) as e:
        _save(array_book, delete_row_nums={2})
    assert "자기 셀을 담지 못합니다" in str(e.value), str(e.value)


# ── 7. 좌표 셈 ──────────────────────────────────────────────────────────────
# 여기가 틀리면 멀쩡한 ref 를 흔들거나 깨진 ref 를 놓친다.

@pytest.mark.parametrize("ref, rng, want", [
    ("B4", "B4:B16", True),
    ("B16", "B4:B16", True),
    ("B3", "B4:B16", False),          # 실기 원인 2 의 모양
    ("B17", "B4:B16", False),
    ("C4", "B4:B16", False),
    ("B4", "B4", True),
    ("B5", "B4", False),
    ("B4", "", True),                 # 판단 불가 — 문제 삼지 않는다
])
def test_ref_containment(ref, rng, want):
    assert _ref_in(ref, rng) is want


@pytest.mark.parametrize("refs, want", [
    (["B4"], "B4"),
    (["B4", "B16"], "B4:B16"),
    (["B16", "B4"], "B4:B16"),
    (["B4", "C9", "A2"], "A2:C9"),
    ([], ""),
])
def test_bounding_ref(refs, want):
    assert _bounding_ref(refs) == want

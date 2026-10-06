# -*- coding: utf-8 -*-
"""저장한 파일이 **Excel 이 열 수 있는 패키지**인가 — 값이 아니라 구조를 본다.

실기에서 드러난 일: 40_Build 에 저장한 파일을 열 때마다 Excel 이 "내용에 문제가
있습니다. 복구하시겠습니까?" 창을 띄웠다. 값은 전부 맞았다 — 그래서 기존 저장 검증
(덮어쓴 셀 + 시트 전체 전수 대조)은 그 파일을 **그대로 통과시켰다**.

원인은 하나였다. 저장은 `xl/calcChain.xml`(수식 계산 순서 캐시)을 항상 빼는데 — 값을
고치면 무효가 되고 Excel 이 다시 만드니 맞는 처리다 — 그것을 **가리키는 참조 둘을
남겼다**:

    [Content_Types].xml          <Override PartName="/xl/calcChain.xml" .../>
    xl/_rels/workbook.xml.rels   rIdN → calcChain.xml

OPC 규약상 선언·참조된 부품은 패키지 안에 있어야 한다. 없으면 Excel 은 복구 창을
띄운다. 실제로 5개 파일(.xlsx 3 / .xlsm 2)이 이 상태로 나갔다.

그래서 이 파일이 고정하는 것 둘:
  1. 뺀 부품을 가리키는 참조가 **남지 않는다**(그리고 멀쩡한 참조는 **지우지 않는다**).
  2. 혹시 남으면 **저장 검증이 막는다** — 값만 보는 검증이 이 사고를 놓쳤으므로,
     구조 검사가 실제로 동작하는지를 테스트가 직접 확인한다(아래 '가드가 살아 있나').
"""
import os
import zipfile

import openpyxl
import pytest

from excelmerge.loaders import clear_values_cache, load_values_any
from excelmerge.xlsx_writer import (_package_mismatches, _rel_target_part,
                                    _write_patches_to_file)

REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
ROWS = [["TID", "NAME", "VALUE"],
        ["1001", "칼", "10"],
        ["1002", "방패", "20"]]


# ── 준비물 ───────────────────────────────────────────────────────────────────

def _add_parts(path, parts: dict, overrides=(), wb_rels=()):
    """만들어 둔 xlsx 에 부품·선언·관계를 끼워 넣는다(실제 파일 모양 재현)."""
    tmp = str(path) + ".build"
    with zipfile.ZipFile(str(path)) as zin, \
         zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in zin.namelist():
            data = zin.read(n)
            if n == "[Content_Types].xml" and overrides:
                data = data.replace(b"</Types>",
                                    "".join(overrides).encode() + b"</Types>")
            elif n == "xl/_rels/workbook.xml.rels" and wb_rels:
                data = data.replace(b"</Relationships>",
                                    "".join(wb_rels).encode() + b"</Relationships>")
            zout.writestr(n, data)
        for n, blob in parts.items():
            zout.writestr(n, blob)
    os.replace(tmp, str(path))


def _rel(rid, type_, target, mode=""):
    mode = f' TargetMode="{mode}"' if mode else ""
    return (f'<Relationship Id="{rid}" '
            f'Type="http://schemas.openxmlformats.org/officeDocument/2006/'
            f'relationships/{type_}" Target="{target}"{mode}/>')


@pytest.fixture
def book(tmp_path):
    """calcChain 이 들어 있는 xlsx — 저장 전 상태가 **깨끗해야** 의미가 있다."""
    p = tmp_path / "a.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for r in ROWS:
        ws.append(r)
    wb.save(str(p))

    _add_parts(
        p,
        parts={"xl/calcChain.xml":
               b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
               b'<calcChain xmlns="http://schemas.openxmlformats.org/'
               b'spreadsheetml/2006/main"><c r="C2" i="1"/></calcChain>'},
        overrides=['<Override PartName="/xl/calcChain.xml" ContentType='
                   '"application/vnd.openxmlformats-officedocument.'
                   'spreadsheetml.calcChain+xml"/>'],
        wb_rels=[_rel("rId900", "calcChain", "calcChain.xml")],
    )
    assert _package_mismatches(str(p)) == [], "전제: 저장 전 패키지는 깨끗하다"
    assert "xl/calcChain.xml" in _names(p), "전제: calcChain 이 들어 있다"
    clear_values_cache()
    return p


def _names(path):
    with zipfile.ZipFile(str(path)) as z:
        return set(z.namelist())


def _part(path, name):
    with zipfile.ZipFile(str(path)) as z:
        return z.read(name).decode("utf-8")


def _save(path, patches=None):
    clear_values_cache()
    _write_patches_to_file(str(path), patches or {"B2": "검"})


# ── 1. 뺀 부품을 가리키는 참조가 남지 않는다 ────────────────────────────────

def test_the_saved_package_has_no_dangling_reference(book):
    """이 테스트가 그 복구 창이다 — 실패하면 Excel 이 복구를 물어본다."""
    _save(book)
    assert _package_mismatches(str(book)) == []


def test_calcchain_itself_is_dropped(book):
    """빼는 것 자체는 맞다 — 값을 고치면 계산 순서 캐시는 무효다."""
    _save(book)
    assert "xl/calcChain.xml" not in _names(book)


def test_the_content_types_override_is_removed(book):
    _save(book)
    assert "calcChain" not in _part(book, "[Content_Types].xml")


def test_the_workbook_relationship_is_removed(book):
    _save(book)
    assert "calcChain" not in _part(book, "xl/_rels/workbook.xml.rels")


def test_the_values_still_land(book):
    """구조를 고치느라 값 쓰기가 망가지면 안 된다."""
    _save(book, {"B2": "검"})
    assert load_values_any(str(book))[1][1] == "검"


# ── 2. 멀쩡한 참조는 지우지 않는다 ───────────────────────────────────────────
# 한쪽으로 치우친 수정은 반대쪽으로 넘어지기 쉽다. '없는 부품을 가리키는 것만' 지운다.

def test_the_other_overrides_survive(book):
    before = _part(book, "[Content_Types].xml")
    _save(book)
    after = _part(book, "[Content_Types].xml")
    for part in ("/xl/workbook.xml", "/xl/worksheets/sheet1.xml",
                 "/xl/styles.xml", "/docProps/core.xml"):
        assert part in before and part in after, f"{part} 선언이 사라졌다"


def test_the_other_relationships_survive(book):
    _save(book)
    rels = _part(book, "xl/_rels/workbook.xml.rels")
    for t in ("worksheets/sheet1.xml", "styles.xml", "theme/theme1.xml"):
        assert t in rels, f"{t} 관계가 사라졌다"


def test_every_remaining_part_is_still_declared(book):
    """반대 방향도 본다 — 부품은 있는데 선언이 사라지면 그것도 복구 대상이다."""
    _save(book)
    ct = _part(book, "[Content_Types].xml")
    for n in _names(book):
        if n.endswith(".rels") or n == "[Content_Types].xml":
            continue
        ext = n.rsplit(".", 1)[-1]
        assert f'/{n}"' in ct or f'Extension="{ext}"' in ct, f"{n} 선언이 없다"


def test_an_external_relationship_is_left_alone(tmp_path):
    """TargetMode=External 은 패키지 밖을 가리킨다 — 경로로 풀어서 보면 안 된다."""
    p = tmp_path / "ext.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(ROWS[0])
    wb.save(str(p))
    _add_parts(p, parts={}, wb_rels=[
        _rel("rId901", "hyperlink", "https://example.com/x", mode="External")])
    _save(p, {"A1": "TID"})
    assert "https://example.com/x" in _part(p, "xl/_rels/workbook.xml.rels")


def test_a_book_without_calcchain_keeps_its_package(tmp_path):
    """뺄 게 없으면 선언·관계를 건드리지 않는다."""
    p = tmp_path / "plain.xlsx"
    wb = openpyxl.Workbook()
    for r in ROWS:
        wb.active.append(r)
    wb.save(str(p))
    before = (_part(p, "[Content_Types].xml"),
              _part(p, "xl/_rels/workbook.xml.rels"))
    _save(p)
    assert (_part(p, "[Content_Types].xml"),
            _part(p, "xl/_rels/workbook.xml.rels")) == before


# ── 3. 가드가 살아 있나 ──────────────────────────────────────────────────────
# 값만 보는 검증이 이 사고를 통과시켰다. 구조 검사를 **넣었다**고 끝이 아니라,
# 그게 실제로 막는지를 봐야 한다 — 버그를 되살려 놓고 저장을 시도한다.

def test_the_verifier_blocks_a_dangling_reference(book, monkeypatch):
    import excelmerge.xlsx_writer as W
    monkeypatch.setattr(W, "_strip_dangling_refs",
                        lambda name, data, dropped: data)   # 버그 되살리기
    before = load_values_any(str(book))[1][1]
    with pytest.raises(ValueError) as e:
        _save(book, {"B2": "검"})
    assert "xl/calcChain.xml" in str(e.value), str(e.value)

    clear_values_cache()
    assert load_values_any(str(book))[1][1] == before, "원본이 덮였다"
    assert "xl/calcChain.xml" in _names(book), "원본이 손상됐다"
    assert not os.path.exists(str(book) + ".tmp_merge"), "임시 파일이 남았다"


def test_the_verifier_names_both_references(book, monkeypatch):
    """어디가 문제인지 말해 주지 않으면 다음 사람이 또 추적부터 시작한다."""
    import excelmerge.xlsx_writer as W
    monkeypatch.setattr(W, "_strip_dangling_refs",
                        lambda name, data, dropped: data)
    with pytest.raises(ValueError) as e:
        _save(book)
    msg = str(e.value)
    assert "[Content_Types].xml" in msg and "workbook.xml.rels" in msg, msg


def test_a_clean_package_reports_nothing(book):
    assert _package_mismatches(str(book)) == []


# ── 4. 참조 경로 해석 ────────────────────────────────────────────────────────
# 여기가 틀리면 둘 중 하나다 — 멀쩡한 관계를 지우거나, 깨진 관계를 놓치거나.

@pytest.mark.parametrize("rels_path, target, want", [
    ("xl/_rels/workbook.xml.rels", "calcChain.xml", "xl/calcChain.xml"),
    ("xl/_rels/workbook.xml.rels", "worksheets/sheet1.xml",
     "xl/worksheets/sheet1.xml"),
    ("xl/worksheets/_rels/sheet1.xml.rels",
     "../printerSettings/printerSettings1.bin",
     "xl/printerSettings/printerSettings1.bin"),
    ("_rels/.rels", "xl/workbook.xml", "xl/workbook.xml"),
    ("_rels/.rels", "docProps/core.xml", "docProps/core.xml"),
    ("xl/_rels/workbook.xml.rels", "/xl/styles.xml", "xl/styles.xml"),
])
def test_relationship_targets_resolve(rels_path, target, want):
    el = _fake_rel(target)
    assert _rel_target_part(rels_path, el) == want


@pytest.mark.parametrize("target, mode", [
    ("https://example.com/x", "External"),
    ("file:///C:/x.xlsx", "External"),
    ("https://example.com/y", ""),     # 모드가 빠져 있어도 URL 은 부품이 아니다
    ("", ""),
])
def test_targets_outside_the_package_are_none(target, mode):
    assert _rel_target_part("xl/_rels/workbook.xml.rels",
                            _fake_rel(target, mode)) is None


def _fake_rel(target, mode=""):
    from lxml import etree
    xml = _rel("rId1", "x", target, mode)
    return etree.fromstring(
        f'<Relationships xmlns="{REL_NS}">{xml}</Relationships>'.encode())[0]

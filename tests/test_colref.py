# -*- coding: utf-8 -*-
"""열 번호 ↔ 열 문자 변환이 openpyxl 과 **완전히 같은가**.

`from openpyxl.utils import get_column_letter` 한 줄이 openpyxl 전체를 끌어온다
(실측: 모듈 430개, 1.14초). 쓰는 건 변환 함수 둘뿐이라 지역 구현으로 바꿨다 —
프로그램 켤 때마다 1초 넘게 기다리던 원인이었다.

이 좌표는 **저장할 셀 자리**를 정하는 데 쓰인다. 다르게 동작하면 엉뚱한 칸에 쓰고,
그건 값 검증을 통과한다(의도한 자리에 의도한 값을 썼으니까). 그래서 전수로 대 본다.
"""
import pytest

from excelmerge.colref import column_index_from_string, get_column_letter

# 대조군은 openpyxl 원본. 느리지만 테스트에서만 쓴다.
from openpyxl.utils import column_index_from_string as op_index
from openpyxl.utils import get_column_letter as op_letter

MAX_COL = 18278          # 'ZZZ'


def test_every_column_letter_matches_openpyxl():
    """1~18278 전부 — 한 칸이라도 어긋나면 저장이 엉뚱한 열로 간다."""
    bad = [i for i in range(1, MAX_COL + 1) if get_column_letter(i) != op_letter(i)]
    assert not bad, f"어긋난 열 번호 {len(bad)}개, 예: {bad[:5]}"


def test_every_column_index_matches_openpyxl():
    bad = []
    for i in range(1, MAX_COL + 1):
        s = op_letter(i)
        if column_index_from_string(s) != op_index(s):
            bad.append(s)
    assert not bad, f"어긋난 열 문자 {len(bad)}개, 예: {bad[:5]}"


def test_round_trip():
    for i in range(1, MAX_COL + 1):
        assert column_index_from_string(get_column_letter(i)) == i


@pytest.mark.parametrize("idx, want", [
    (1, "A"), (26, "Z"), (27, "AA"), (52, "AZ"), (53, "BA"),
    (702, "ZZ"), (703, "AAA"), (16384, "XFD"), (18278, "ZZZ"),
])
def test_known_boundaries(idx, want):
    """손으로 짚어 둔 경계 — 26 의 거듭제곱 언저리에서 틀리기 쉽다."""
    assert get_column_letter(idx) == want
    assert column_index_from_string(want) == idx


# ── 범위 밖·잘못된 입력도 openpyxl 과 같이 거절한다 ────────────────────────
# 조용히 다른 값을 돌려주면 엉뚱한 칸에 쓴다. 터지는 편이 낫다.

@pytest.mark.parametrize("idx", [0, -1, MAX_COL + 1, 100000])
def test_out_of_range_index_raises_like_openpyxl(idx):
    with pytest.raises(ValueError):
        get_column_letter(idx)
    with pytest.raises(ValueError):
        op_letter(idx)


@pytest.mark.parametrize("col", ["", "AAAA", "A1", "1", "@", "A-", "가", " A"])
def test_bad_column_name_raises_like_openpyxl(col):
    with pytest.raises(ValueError):
        column_index_from_string(col)
    with pytest.raises(ValueError):
        op_index(col)


@pytest.mark.parametrize("col, want", [("a", 1), ("zz", 702), ("xFd", 16384)])
def test_lowercase_is_accepted(col, want):
    """셀 참조를 소문자로 쓴 파일이 있다 — openpyxl 과 같이 받아 준다."""
    assert column_index_from_string(col) == want == op_index(col)


# ── 시작 경로에 openpyxl 이 없어야 한다 ────────────────────────────────────

def test_colref_does_not_import_openpyxl():
    """이 모듈만 불러서는 openpyxl 이 딸려오면 안 된다 — 그러면 고친 의미가 없다."""
    import subprocess
    import sys
    out = subprocess.run(
        [sys.executable, "-c",
         "import sys; import excelmerge.colref; "
         "print('openpyxl' in sys.modules)"],
        capture_output=True, text=True, cwd=__import__("os").path.dirname(
            __import__("os").path.dirname(__file__)))
    assert out.stdout.strip() == "False", out.stdout + out.stderr

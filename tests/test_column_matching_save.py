# -*- coding: utf-8 -*-
"""열 매칭 3단계 — 저장이 **대상 파일의 제 열**에 쓴다.

행은 row_meta 로 화면 행 → 파일 행을 옮겨 왔다. 열도 같은 변환이 필요하다. 이걸
빼먹으면 A 의 값이 B 의 엉뚱한 열에 들어가고, 셀 단위 검증은 '의도한 자리에 의도한
값' 이라며 통과한다(시트 전수 대조가 마지막 그물이다).

여기서 고정하는 것:
  1. 화면 열 → 그 side 의 파일 열로 옮겨 쓴다(덮어쓰기·신규 행·서식 소스 모두).
  2. 대상 파일에 없는 열은 쓰지 않는다 — 열 삽입은 지원하지 않는다.
  3. col_meta 가 없으면 좌표가 예전과 **완전히 같다**.
"""
import openpyxl
import pytest

from excelmerge.constants import DIR_A2B, DIR_B2A, STATUS_ADDED, STATUS_MODIFIED
from excelmerge.diff_engine import compute_diff, match_columns
from excelmerge.loaders import clear_values_cache, load_values_any
from excelmerge.merge_build import build_side_patches
from excelmerge.staging import stageable_cells
from excelmerge.xlsx_writer import _promote_empty_cols_to_delete, _write_patches_to_file

# A: ID | NAME | VALUE       B: ID | GRADE | NAME | VALUE
A_DATA = [["ID", "NAME", "VALUE"], ["k1", "칼", "10"], ["k2", "방패", "20"]]
B_DATA = [["ID", "GRADE", "NAME", "VALUE"],
          ["k1", "A", "검", "10"], ["k2", "S", "방패", "99"]]


def _prep(a=A_DATA, b=B_DATA, key_col=0):
    cm = match_columns(a, b)
    m, meta = compute_diff(a, b, key_col=key_col, key_row=0, col_meta=cm)
    return m, meta, cm


def _stage_all(m, direction, rows=None):
    cells = set()
    for r in range(len(m)):
        if rows is not None and r not in rows:
            continue
        cells |= stageable_cells(m, {(r, c) for c in range(len(m[r]))})
    return {c: direction for c in cells}


# ── 좌표 변환 ────────────────────────────────────────────────────────────────

def test_a_to_b_writes_into_bs_own_columns():
    """화면 2열(NAME)은 A 의 B열이지만 B 파일에서는 C열이다."""
    m, meta, cm = _prep()
    assert cm == [(0, 0), (None, 1), (1, 2), (2, 3)]
    staged = _stage_all(m, DIR_A2B, rows={1})
    patches, ins, _style = build_side_patches(DIR_A2B, m, meta, staged, cm)
    assert patches == {"C2": "칼"}, patches       # NAME: 화면 2 → B 파일 C열
    assert ins == {}


def test_without_col_meta_the_old_coordinates_are_used():
    """예전 호출(맞추지 않은 매트릭스 + col_meta 없음)은 한 글자도 달라지면 안 된다."""
    m, meta = compute_diff(A_DATA, B_DATA, key_col=0, key_row=0)   # 위치 기준
    staged = _stage_all(m, DIR_A2B, rows={1})
    patches, _ins, _s = build_side_patches(DIR_A2B, m, meta, staged)
    assert set(patches) == {"B2", "C2", "D2"}, patches


def test_b_to_a_writes_into_as_own_columns():
    m, meta, cm = _prep()
    staged = _stage_all(m, DIR_B2A, rows={1})
    patches, _ins, _s = build_side_patches(DIR_B2A, m, meta, staged, cm)
    assert patches == {"B2": "검"}, patches        # NAME: 화면 2 → A 파일 B열


def test_the_style_source_is_mapped_too():
    """서식을 가져올 좌표도 소스 파일 기준이어야 한다 — 안 그러면 옆 열 서식이 붙는다."""
    m, meta, cm = _prep()
    staged = _stage_all(m, DIR_A2B, rows={1})
    _p, _i, style = build_side_patches(DIR_A2B, m, meta, staged, cm)
    assert style == {"C2": "B2"}, style            # 대상 B 의 C2 ← 소스 A 의 B2


def test_a_column_missing_on_the_target_is_skipped():
    """B→A 에서 GRADE 는 A 에 없다 — 쓸 자리가 없으니 건너뛴다(열 삽입 미지원)."""
    m, meta, cm = _prep()
    staged = {(1, 1): DIR_B2A}                     # GRADE 만 준비
    patches, ins, _s = build_side_patches(DIR_B2A, m, meta, staged, cm)
    assert patches == {} and ins == {}


def test_a_new_row_uses_target_columns():
    """B 에만 있는 행을 B→A 로 넣을 때, 각 값이 A 의 제 열로 가야 한다."""
    a = [["ID", "NAME"], ["k1", "칼"]]
    b = [["ID", "GRADE", "NAME"], ["k1", "A", "칼"], ["k9", "S", "창"]]
    cm = match_columns(a, b)
    m, meta = compute_diff(a, b, key_col=0, key_row=0, col_meta=cm)
    r = next(i for i in range(len(m)) if meta[i] == (None, 2))
    staged = _stage_all(m, DIR_B2A, rows={r})
    _p, ins, _s = build_side_patches(DIR_B2A, m, meta, staged, cm)
    cols = {c: v for c, v, _src in ins[r]}
    assert cols == {0: "k9", 1: "창"}, cols        # GRADE 는 A 에 없어 빠진다


# ── 실제 파일에 써 본다 ─────────────────────────────────────────────────────

def _xlsx(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    wb.save(str(path))
    return str(path)


def _save(target, m, meta, cm, staged, direction, src):
    patches, ins, style = build_side_patches(direction, m, meta, staged, cm)
    patches, dels, delcols = _promote_empty_cols_to_delete(patches, set(), target)
    style = {k: v for k, v in style.items() if k in patches}
    _write_patches_to_file(target, patches, list(ins.values()), dels, delcols,
                           None, src, None, style)
    clear_values_cache()
    return load_values_any(target)


def test_merging_into_a_shifted_file_lands_in_the_right_column(tmp_path):
    """가장 중요한 한 줄 — 열이 어긋난 파일에 A→B 병합해도 값이 제자리에 들어간다."""
    pa = _xlsx(tmp_path / "a.xlsx", A_DATA)
    pb = _xlsx(tmp_path / "b.xlsx", B_DATA)
    m, meta, cm = _prep()
    staged = _stage_all(m, DIR_A2B, rows={1, 2})
    rows = _save(pb, m, meta, cm, staged, DIR_A2B, pa)
    assert rows[0] == ["ID", "GRADE", "NAME", "VALUE"], "헤더가 망가졌다"
    assert rows[1] == ["k1", "A", "칼", "10"], rows[1]    # NAME 만 A 값으로, GRADE 보존
    assert rows[2] == ["k2", "S", "방패", "20"], rows[2]  # VALUE 만 A 값으로


def test_the_untouched_one_sided_column_survives(tmp_path):
    """GRADE 는 A 에 없다 — 병합해도 B 의 값이 지워지면 안 된다."""
    pa = _xlsx(tmp_path / "a.xlsx", A_DATA)
    pb = _xlsx(tmp_path / "b.xlsx", B_DATA)
    m, meta, cm = _prep()
    rows = _save(pb, m, meta, cm, _stage_all(m, DIR_A2B), DIR_A2B, pa)
    assert [r[1] for r in rows[1:]] == ["A", "S"], rows


def test_the_whole_file_check_would_catch_a_wrong_mapping(tmp_path, monkeypatch):
    """좌표를 한 칸 밀어 쓰면 전수 대조가 막아야 한다 — 이 변환의 마지막 그물."""
    import excelmerge.xlsx_writer as xw
    from openpyxl.utils import column_index_from_string, get_column_letter
    pa = _xlsx(tmp_path / "a.xlsx", A_DATA)
    pb = _xlsx(tmp_path / "b.xlsx", B_DATA)
    before = open(pb, "rb").read()
    real = xw._patch_sheet_xml

    def shift(data, patches, *a, **k):
        moved = {}
        for ref, v in patches.items():
            mm = xw._COL_RE.match(ref)
            col = column_index_from_string(mm.group(1))
            moved[f"{get_column_letter(col + 1)}{mm.group(2)}"] = v
        return real(data, moved, *a, **k)

    monkeypatch.setattr(xw, "_patch_sheet_xml", shift)
    m, meta, cm = _prep()
    with pytest.raises(ValueError, match="저장 검증 실패"):
        _save(pb, m, meta, cm, _stage_all(m, DIR_A2B, rows={1}), DIR_A2B, pa)
    assert open(pb, "rb").read() == before, "검증 실패인데 원본이 바뀌었다"


def test_deleting_a_b_only_row_still_works_with_matching(tmp_path):
    """행 삭제(한쪽에만 있는 행을 A→B 로 지우기)가 열 매칭과 함께서도 동작해야 한다."""
    a = [["ID", "NAME"], ["k1", "칼"]]
    b = [["ID", "GRADE", "NAME"], ["k1", "A", "칼"], ["k9", "S", "창"]]
    pa = _xlsx(tmp_path / "a.xlsx", a)
    pb = _xlsx(tmp_path / "b.xlsx", b)
    cm = match_columns(a, b)
    m, meta = compute_diff(a, b, key_col=0, key_row=0, col_meta=cm)
    r = next(i for i in range(len(m)) if meta[i] == (None, 2))   # k9 = B 의 3행
    rows = _save(pb, m, meta, cm, _stage_all(m, DIR_A2B, rows={r}), DIR_A2B, pa)
    assert [x[0] for x in rows] == ["ID", "k1"], f"빈 행이 남았거나 잘못 지웠다: {rows}"
